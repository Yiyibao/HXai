"""Independent readers validate Android-produced files; Office rendering is an additional check."""
import hashlib
import json
import pathlib
import unicodedata
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader
import pypdfium2

qa = pathlib.Path(__file__).resolve().parents[1] / 'qa/material'
results = []
for path in sorted(qa.iterdir()):
    if path.name == 'real-model.md' or not (path.name.startswith('sample.') or path.name.startswith('real-') or path.name == 'multipage.pdf'):
        continue
    ext = path.suffix.lower()
    if ext == '.docx':
        doc = Document(path)
        text = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    elif ext == '.xlsx':
        book = load_workbook(path, data_only=False)
        text = '\n'.join(str(c.value or '') for s in book for row in s for c in row)
        assert all(c.data_type != 'f' for s in book for row in s for c in row), 'No executable formula is promised'
    elif ext == '.pptx':
        deck = Presentation(path)
        text = '\n'.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
    elif ext == '.md':
        text = path.read_text(encoding='utf-8')
    elif ext == '.pdf':
        pdf = PdfReader(path)
        text = '\n'.join(p.extract_text() for p in pdf.pages)
        # Only CJK radical glyph mappings are normalized; superscript math remains unchanged.
        text = ''.join(unicodedata.normalize('NFKC', c) if 0x2F00 <= ord(c) <= 0x2FD5 else c for c in text)
        render = pypdfium2.PdfDocument(str(path))
        for i, page in enumerate(render):
            page.render(scale=1.25).to_pil().save(qa / f'{path.stem}-page-{i+1}.png')
    else:
        continue
    assert '中文' in text, f'{path.name}: actual Chinese content missing'
    results.append({'file': path.name, 'format': ext[1:], 'bytes': path.stat().st_size, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'independent_read': 'PASS'})
(qa / 'independent-results.json').write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'files': len(results), 'independent_read': 'PASS'}))
