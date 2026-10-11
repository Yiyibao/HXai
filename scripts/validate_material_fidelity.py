"""Independent checks of Android fixtures, not the app's own parser."""
from pathlib import Path
import json
import hashlib
import unicodedata
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

qa = Path(__file__).resolve().parents[1] / 'qa/material'
results = {}
word = Document(qa / 'fidelity.docx')
assert len(word.inline_shapes) == 1
assert word.tables[0].cell(1, 0).text == '样例'
with ZipFile(qa / 'fidelity.docx') as archive:
    root = ET.fromstring(archive.read('word/document.xml'))
    ns = {'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    assert root.find('.//m:f/m:num/m:r/m:t', ns).text == 'x+1'
    assert root.find('.//m:f/m:den/m:r/m:t', ns).text == '2'
    assert root.find('.//m:sSup/m:sup/m:r/m:t', ns).text == '3'
results['word'] = {'image_count': len(word.inline_shapes), 'editable_omml_fraction_and_power': 'PASS', 'table': 'PASS'}
book = load_workbook(qa / 'fidelity.xlsx', data_only=False)
assert book.sheetnames == ['成绩与编号']
sheet = book.active
assert sheet['A2'].value == '0012' and sheet['A2'].data_type == 's'
assert sheet['B2'].value == 12.5 and sheet['B2'].data_type == 'n'
assert sheet['C2'].value is True and sheet['C2'].data_type == 'b'
assert sheet['D2'].value == '=1+1' and sheet['D2'].data_type == 's'
assert sheet['E2'].value is None
assert sheet['A2'].alignment.wrap_text and sheet.sheet_properties.pageSetUpPr.fitToPage
assert sheet['A2'].border.left.style == 'thin'
order = load_workbook(qa / 'sheet-order.xlsx')
assert order.sheetnames == ['先读的工作表', '成绩与编号'] and order.worksheets[0]['A1'].value == 99
results['excel'] = {'names_and_relationship_order': 'PASS', 'number_boolean_string_null': 'PASS', 'formula_string': 'PASS', 'wrap_and_print_width': 'PASS'}
pdf = PdfReader(qa / 'table-long.pdf')
def normalized(text):
    return ''.join(unicodedata.normalize('NFKC', c) if 0x2f00 <= ord(c) <= 0x2fd5 else c for c in text)
text = ''.join(normalized('\n'.join(page.extract_text() or '' for page in pdf.pages)).split())
assert '第160行' in text and '第80行' in text
table_pages = [i + 1 for i, page in enumerate(pdf.pages) if '序号' in normalized(page.extract_text() or '')]
assert len(table_pages) > 5
assert table_pages == list(range(min(table_pages), len(pdf.pages) + 1))
results['pdf'] = {'pages': len(pdf.pages), 'table_pages': table_pages, 'tail_row_and_repeated_header': 'PASS', 'visual_review': 'separate PNG inspection required'}
long = (qa / 'long.md').read_text(encoding='utf-8')
assert len(long) > 24000 and long.endswith('TAIL_2000')
results['long_text'] = {'characters': len(long), 'tail': 'PASS'}
if (qa / 'scan-input.pdf').exists() and (qa / 'scan-output.md').exists():
    scan = PdfReader(qa / 'scan-input.pdf')
    assert all(not page.extract_text().strip() for page in scan.pages)
    answer = (qa / 'scan-output.md').read_text(encoding='utf-8')
    assert '7391' in answer and 'x' in answer and '2' in answer
    results['scan'] = {'no_text_layer': 'PASS', 'model_output_identifier_and_fraction': 'PASS'}
if (qa / 'word-vision-output.docx').exists():
    output = Document(qa / 'word-vision-output.docx')
    assert len(output.inline_shapes) == 1
    with ZipFile(qa / 'word-vision-output.docx') as archive, ZipFile(qa / 'fidelity.docx') as source:
        root = ET.fromstring(archive.read('word/document.xml'))
        assert root.find('.//m:f', ns) is not None and root.find('.//m:sSup', ns) is not None
        assert archive.read('word/media/image1.png') == source.read('word/media/image1.png')
    results['word_model'] = {'actual_digit_8642': 'PASS by Android assertion', 'image_bytes_preserved': 'PASS', 'editable_omml_preserved': 'PASS'}
results['sha256'] = {name: hashlib.sha256((qa / name).read_bytes()).hexdigest() for name in ['fidelity.docx', 'fidelity.xlsx', 'table-long.pdf', 'long.md']}
(qa / 'fidelity-independent.json').write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(results, ensure_ascii=False, indent=2))
