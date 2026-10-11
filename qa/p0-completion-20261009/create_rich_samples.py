from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PptInches
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image, ImageDraw

root = Path(__file__).parent / 'samples'
root.mkdir(exist_ok=True)
image = Image.new('RGB', (400, 200), 'white')
draw = ImageDraw.Draw(image)
draw.rectangle((30, 30, 170, 160), fill='navy')
draw.ellipse((210, 30, 360, 180), fill='orange')
image.save(root / 'synthetic.png')
doc = Document()
p = doc.add_paragraph()
p.add_run('Bold source').bold = True
p.add_run(' and plain source').italic = True
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = 'Item'; table.cell(0, 1).text = 'Quantity'
table.cell(1, 0).text = 'Synthetic'; table.cell(1, 1).text = '21'
doc.add_picture(str(root / 'synthetic.png'), width=Inches(2))
eq = OxmlElement('m:oMath'); run = OxmlElement('m:r'); text = OxmlElement('m:t'); text.text = 'x² + y² = z²'
run.append(text); eq.append(run); doc.add_paragraph()._p.append(eq)
doc.sections[0].header.paragraphs[0].text = 'Synthetic header'
doc.save(root / 'rich-source.docx')
prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[5]); slide.shapes.title.text = 'Synthetic slide title'
slide.shapes.add_picture(str(root / 'synthetic.png'), PptInches(1), PptInches(2), width=PptInches(3))
box = slide.shapes.add_textbox(PptInches(5), PptInches(2), PptInches(3), PptInches(1)); box.text = 'Keep adjacent text'
prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = 'Untouched second slide'
prs.save(root / 'rich-source.pptx')
wb = Workbook(); ws = wb.active; ws.title = 'Data'
ws.append(['Count', 'Triple', 'Code']); ws.append([2, '=A2*3', '001'])
ws['A2'].font = Font(bold=True, color='000080'); ws.add_image(__import__('openpyxl.drawing.image', fromlist=['Image']).Image(root / 'synthetic.png'), 'A5')
wb.save(root / 'rich-source.xlsx')
print('Created four non-private synthetic fixture files')
