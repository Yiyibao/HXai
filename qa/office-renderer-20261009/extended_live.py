"""Synthetic print-layout boundaries, cancellation tombstone and unchanged MinerU API."""
from pathlib import Path
import json,uuid,hashlib,time,urllib.request
from docx import Document
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font
from pptx import Presentation
from pptx.util import Pt
from pypdf import PdfReader

root=Path(__file__).resolve().parents[2]; evidence=Path(__file__).parent/'extended-output'; evidence.mkdir(exist_ok=True)
secret=json.loads((root/'.local/document-service/credentials/office-test-20261009.json').read_text())
headers={'Authorization':'Bearer '+secret['token']}
def call(path,data=None,extra=None,method=None):
    return urllib.request.urlopen(urllib.request.Request(secret['baseUrl']+path,data=data,headers={**headers,**(extra or {})},method=method),timeout=120)
def wait(task):
    deadline=time.monotonic()+960
    while task['state'] not in {'SUCCEEDED','FAILED','CANCELLED','EXPIRED'}:
        assert time.monotonic()<deadline
        time.sleep(2); task=json.load(call('/api/v1/jobs/'+task['taskId']))
    return task
def submit(file):
    return json.load(call('/api/v1/conversions?extension='+file.suffix[1:],file.read_bytes(),{'Idempotency-Key':str(uuid.uuid4()),'X-Source-Sha256':hashlib.sha256(file.read_bytes()).hexdigest()}))

doc=Document(); style=doc.styles['Normal']; style.font.name='Noto Sans CJK SC'; style._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
doc.sections[0].header.paragraphs[0].text='Synthetic repeated header'
doc.sections[0].footer.paragraphs[0].text='Synthetic repeated footer'
doc.add_heading('中文排版与分页验收',0); doc.add_paragraph('第一页中文、表格与原始分页。')
table=doc.add_table(rows=2,cols=2); table.cell(0,0).text='项目'; table.cell(0,1).text='数值'; table.cell(1,0).text='中文'; table.cell(1,1).text='42'
doc.add_page_break(); doc.add_heading('第二页中文',1); doc.add_paragraph('Explicit second page.'); doc.save(evidence/'layout.docx')
wb=Workbook(); ws=wb.active; ws.title='Visible'; ws.append(['中文项目','Recalculated']); ws.append(['合计','=4*3'])
for row in ws:
    for cell in row: cell.font=Font(name='Noto Sans CJK SC')
ws.column_dimensions['A'].width=24; ws.column_dimensions['B'].width=24
ws['Z100']='OUTSIDE_PRINT_AREA'; ws.print_area='A1:B2'; ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=ws.PAPERSIZE_A4
ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=1; ws.sheet_properties.pageSetUpPr.fitToPage=True
hidden=wb.create_sheet('Hidden'); hidden['A1']='HIDDEN_SHEET_SECRET'; hidden.sheet_state='hidden'; wb.save(evidence/'layout.xlsx')
prs=Presentation()
for title in ['VISIBLE_FIRST','HIDDEN_SLIDE_SECRET','VISIBLE_LAST']:
    slide=prs.slides.add_slide(prs.slide_layouts[5]); slide.shapes.title.text=title
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size=Pt(32)
    slide.notes_slide.notes_text_frame.text='NOTES_SECRET'
prs.slides[1]._element.set('show','0')
ids=prs.slides._sldIdLst; first=ids[0]; ids.remove(first); ids.append(first)  # Actual source order now hidden,last,first.
prs.save(evidence/'layout.pptx')
bad=Workbook(); bad.active['A1']='=1/0'; bad.save(evidence/'formula-error.xlsx')
records=[]
for extension in ['docx','xlsx','pptx']:
    source=evidence/('layout.'+extension); digest=hashlib.sha256(source.read_bytes()).hexdigest(); task=submit(source); ident=task['taskId']
    try:
        task=wait(task); assert task['state']=='SUCCEEDED',task
        manifest=json.load(call('/api/v1/jobs/'+ident+'/manifest')); pdf=call('/api/v1/jobs/'+ident+'/assets/pdf').read()
        assert hashlib.sha256(pdf).hexdigest()==manifest['assets'][0]['sha256']
        target=evidence/(extension+'.pdf'); target.write_bytes(pdf); reader=PdfReader(target); texts=[page.extract_text() for page in reader.pages]; text='\n'.join(texts)
        if extension=='docx':
            assert len(texts)==2 and '第二页中文' in texts[1] and all('Synthetic repeated header' in p and 'Synthetic repeated footer' in p for p in texts)
        elif extension=='xlsx':
            assert len(texts)==1 and '12' in text and '中文项目' in text and 'OUTSIDE_PRINT_AREA' not in text and 'HIDDEN_SHEET_SECRET' not in text
            assert reader.pages[0].mediabox.width>reader.pages[0].mediabox.height
        else:
            assert len(texts)==2 and 'VISIBLE_LAST' in texts[0] and 'VISIBLE_FIRST' in texts[1] and 'HIDDEN_SLIDE_SECRET' not in text and 'NOTES_SECRET' not in text
        assert hashlib.sha256(source.read_bytes()).hexdigest()==digest
        records.append({'format':extension,'manifest':manifest,'text':texts,'originalUnchanged':True}); print(extension,'extended layout passed',flush=True)
    finally: json.load(call('/api/v1/jobs/'+ident,method='DELETE'))
task=submit(evidence/'formula-error.xlsx'); ident=task['taskId']
try:
    task=wait(task); assert task['state']=='FAILED' and task['errorCode']=='formula_recalculation_failed',task
    records.append({'formulaErrorExplicitlyFailed':True}); print('formula error rejected explicitly',flush=True)
finally: json.load(call('/api/v1/jobs/'+ident,method='DELETE'))
key=str(uuid.uuid4()); cancelled=json.load(call('/api/v1/conversion-requests/'+key+'/cancel',b'')); assert cancelled['state']=='CANCELLED'
source=evidence/'layout.docx'; task=json.load(call('/api/v1/conversions?extension=docx',source.read_bytes(),{'Idempotency-Key':key,'X-Source-Sha256':hashlib.sha256(source.read_bytes()).hexdigest()}))
assert task['state']=='CANCELLED'; json.load(call('/api/v1/jobs/'+task['taskId'],method='DELETE')); records.append({'cancelBeforeUploadRemainsCancelled':True})
source=root/'qa/office-renderer-20261009/live-api-output/docx.pdf'; data=source.read_bytes()
task=json.load(call('/api/v1/jobs?extension=pdf&start_page=1&end_page=1',data,{'Idempotency-Key':str(uuid.uuid4()),'X-Source-Sha256':hashlib.sha256(data).hexdigest()})); ident=task['taskId']
try:
    task=wait(task); assert task['state']=='SUCCEEDED',task
    manifest=json.load(call('/api/v1/jobs/'+ident+'/manifest')); assert manifest['processedPages']==[1] and manifest['sourceSha256']==hashlib.sha256(data).hexdigest()
    document=json.load(call('/api/v1/jobs/'+ident+'/assets/document')); assert document['blocks']
    records.append({'existingMinerUParseStillWorks':True,'processedPages':manifest['processedPages']}); print('Existing real MinerU parse passed',flush=True)
finally: json.load(call('/api/v1/jobs/'+ident,method='DELETE'))
(evidence/'results.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
